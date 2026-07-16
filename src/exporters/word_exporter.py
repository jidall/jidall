from __future__ import annotations

from collections import Counter
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from models.assessment import GenerationResult, Question


class WordExporter:
    def export(self, result: GenerationResult, output_dir: Path, output_name: str = "") -> tuple[Path, Path, Path]:
        output_dir.mkdir(parents=True, exist_ok=True)
        questions_path = output_dir / "Questoes.docx"
        answers_path = output_dir / "Gabarito.docx"
        audit_path = output_dir / "Relatorio_Auditoria.docx"
        self._write_questions(result.questions, result, questions_path)
        self._write_answers(result.questions, answers_path)
        self._write_audit(result, audit_path)
        return questions_path, answers_path, audit_path

    def _base_doc(self) -> Document:
        doc = Document()
        style = doc.styles["Normal"]
        style.font.name = "Arial"
        style.font.size = Pt(12)
        return doc

    def _write_questions(self, questions: list[Question], result: GenerationResult, path: Path) -> None:
        doc = self._base_doc()
        doc.add_heading("🎓 Assistente de Provas", level=1)
        doc.add_heading("Questões", level=2)
        for index, question in enumerate(questions, start=1):
            doc.add_heading(f"Questão {index} - {question.kind.title()}", level=2)
            if question.category:
                self._p(doc, "Grupo: " + question.category)
            if question.difficulty:
                self._p(doc, "Dificuldade: " + question.difficulty)
            self._p(doc, "Introdução: " + question.introduction)
            self._p(doc, "Comando: " + question.command)
            for alt in question.alternatives:
                self._p(doc, f"{alt.label}) {alt.text}")
            if question.answer:
                self._p(doc, "Padrão de resposta esperado: " + question.answer)
            self._write_metadata(doc, question, include_answer=False)
        doc.save(path)

    def _write_answers(self, questions: list[Question], path: Path) -> None:
        doc = self._base_doc()
        doc.add_heading("🎓 Assistente de Provas", level=1)
        doc.add_heading("Gabarito", level=2)
        for index, question in enumerate(questions, start=1):
            doc.add_heading(f"Questão {index}", level=2)
            if question.alternatives:
                correct = next((alt.label for alt in question.alternatives if alt.is_correct), "")
                self._p(doc, "Gabarito: " + correct)
                for alt in question.alternatives:
                    self._p(doc, f"Justificativa {alt.label}: {alt.justification}")
            elif question.answer:
                self._p(doc, "Padrão de resposta esperado: " + question.answer)
            self._write_metadata(doc, question, include_answer=True)
        doc.save(path)

    def _write_metadata(self, doc: Document, question: Question, include_answer: bool) -> None:
        for label, value in (("Referência", question.reference), ("Taxonomia de Bloom", question.bloom_taxonomy), ("Competência", question.competence), ("Habilidade", question.skill)):
            if value:
                self._p(doc, f"{label}: {value}")

    def _write_audit(self, result: GenerationResult, path: Path) -> None:
        doc = self._base_doc()
        questions = result.questions
        doc.add_heading("🎓 Assistente de Provas", level=1)
        doc.add_heading("Relatório de Auditoria", level=2)
        self._p(doc, f"Total gerado: {len(questions)}")
        self._p(doc, f"Total aprovado sem alertas: {max(0, len(questions) - len(result.audit_messages))}")
        self._p(doc, f"Total rejeitado/pendente de revisão: {len(result.audit_messages)}")
        self._p(doc, "Distribuição de gabaritos: " + self._answer_distribution(questions))
        self._p(doc, "Distribuição de dificuldade: " + str(Counter((q.difficulty or "não informada") for q in questions)))
        self._p(doc, "Distribuição de Bloom: " + str(Counter((q.bloom_taxonomy or "não informada") for q in questions)))
        doc.add_heading("Documentos e fontes utilizadas", level=2)
        for source in self._sources(questions):
            self._p(doc, source)
        doc.add_heading("Divergências", level=2)
        for item in result.theme_map.divergences or ["Nenhuma divergência automática registrada."]:
            self._p(doc, item)
        doc.add_heading("Falhas e limitações", level=2)
        for item in result.audit_messages or ["Nenhum alerta de validação registrado."]:
            self._p(doc, item)
        doc.save(path)

    def _answer_distribution(self, questions: list[Question]) -> str:
        answers = [next((alt.label for alt in q.alternatives if alt.is_correct), "") for q in questions if q.alternatives]
        return str(Counter(answer for answer in answers if answer))

    def _sources(self, questions: list[Question]) -> list[str]:
        values = set()
        for question in questions:
            for snippet in question.source_snippets:
                values.add(f"{snippet.document_name} | {snippet.aula or 'Aula não identificada'} | {snippet.theme or 'Tema não identificado'} | página {snippet.page_number}")
        return sorted(values) or ["Nenhuma fonte rastreável associada às questões."]

    def _p(self, doc: Document, text: str) -> None:
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
